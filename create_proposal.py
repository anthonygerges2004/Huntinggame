from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm

OUTPUTS = [
    r"C:\Users\antho\senior\CSIS290_Hunting_Game_Project_Proposal.docx",
    r"C:\Users\antho\senior\CSIS290_Hunting_Game_Project_Proposal_UPDATED.docx",
    r"C:\Users\antho\OneDrive\Desktop\CSIS290_Hunting_Game_Project_Proposal.docx",
    r"C:\Users\antho\OneDrive\Documents\CSIS290_Hunting_Game_Project_Proposal.docx",
    r"d:\downloadss\CSIS290_Hunting_Game_Project_Proposal.docx",
    r"C:\senior\CSIS290_Hunting_Game_Project_Proposal.docx",
]

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = "1F3A5F"
LIGHT = "E8EEF6"
GANTT_FILL = "2E6B4F"
GANTT_EMPTY = "F4F6F8"


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_paragraph(doc, text, *, size=12, bold=False, italic=False, align="left", space_after=8, space_before=0, first_line=False):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.15
    if first_line:
        pf.first_line_indent = Cm(1.0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "9AA4B2")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def write_cell(cell, text, *, bold=False, size=10, center=False, color=None, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_border(cell)
    if fill:
        shade_cell(cell, fill)


def heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{number}. {title}")
    set_run_font(run, size=14, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), ACCENT)
    pBdr.append(bottom)
    pPr.append(pBdr)


def subheading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_run_font(run, size=12, bold=True, color=NAVY)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=12)


def set_page(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("CSIS290  |  Project Proposal  |  Wildlife Hunting Game")
    set_run_font(run, size=9, italic=True, color=RGBColor(0x5B, 0x67, 0x75))
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Page ")
    set_run_font(run, size=9, color=RGBColor(0x5B, 0x67, 0x75))
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r1 = fp.add_run()
    r1._r.append(fld)
    r2 = fp.add_run()
    r2._r.append(instr)
    r3 = fp.add_run()
    r3._r.append(fld2)
    set_run_font(r1, size=9, color=RGBColor(0x5B, 0x67, 0x75))
    set_run_font(r2, size=9, color=RGBColor(0x5B, 0x67, 0x75))
    set_run_font(r3, size=9, color=RGBColor(0x5B, 0x67, 0x75))


def add_architecture_table(doc):
    subheading(doc, "Component Interaction")
    table = doc.add_table(rows=9, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer", "Components", "Responsibility"]
    rows = [
        (
            "Presentation",
            "Camera, HUD (health, time of day, inventory), crosshair",
            "Show first-person view, life, day/night, and carried items",
        ),
        (
            "Player systems",
            "PlayerMove, PlayerShoot, PlayerHealth, Character Controller",
            "Walk, look, crouch, shoot, take damage, and die at zero health",
        ),
        (
            "Hunting & wildlife",
            "Animal types (birds, deer, rare large animals), loot drops",
            "Spawn by rarity; scatter/flee on shots; drop meat; legendary rare spawn",
        ),
        (
            "Survival",
            "Campfire, DayNightCycle, NightThreat",
            "Morning/night lighting; night attacks unless a fire is lit",
        ),
        (
            "Gathering",
            "Tree, Pickaxe tool check, Inventory",
            "Chop wood only if the player has a pickaxe",
        ),
        (
            "Trading",
            "Trader camp / trade post",
            "Buy/sell meat, wood, weapons; trader may also hunt the legendary animal",
        ),
        (
            "Special gameplay",
            "Alarm (birds scatter), MeatSpoil, Cooking, LegendaryAnimal, Tracks, Overhunt, RivalTrader, Stealth",
            "Wildlife reacts to shots; meat goes bad unless cooked; rare legendary hunt; tracks; quiet movement",
        ),
        (
            "Engine / input",
            "Unity 6, URP, Input System, Physics, Lighting",
            "Render the world, read input, raycasts, day/night lights",
        ),
    ]
    for i, h in enumerate(headers):
        write_cell(table.rows[0].cells[i], h, bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    for r, row in enumerate(rows, start=1):
        fill = LIGHT if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            write_cell(table.rows[r].cells[c], val, size=10, fill=fill, center=(c == 0))


def add_flow_table(doc):
    subheading(doc, "Core Gameplay Flow")
    table = doc.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    steps = [
        "Explore wildlife area",
        "Hunt / chop wood",
        "Loot / cook meat",
        "Fire & night",
        "Trade at camp",
        "Survive (health)",
    ]
    for i, step in enumerate(steps):
        write_cell(table.rows[0].cells[i], step, bold=True, size=9, center=True, color=RGBColor(255, 255, 255), fill=GANTT_FILL)
        write_cell(table.rows[1].cells[i], "→" if i < 5 else "Continue cycle", size=8, center=True, fill=LIGHT)


def add_gantt(doc):
    add_paragraph(
        doc,
        "The timeline covers the remainder of the Fall 2026 term, from 11 September 2026 to 8 December 2026. "
        "Shaded cells show when a task is active. Hunting, survival, and trading work are split between the two members so several tracks can run in parallel.",
        align="justify",
        space_after=8,
    )
    weeks = [
        "W1 11–17 Sep",
        "W2 18–24 Sep",
        "W3 25 Sep–1 Oct",
        "W4 2–8 Oct",
        "W5 9–15 Oct",
        "W6 16–22 Oct",
        "W7 23–29 Oct",
        "W8 30 Oct–5 Nov",
        "W9 6–12 Nov",
        "W10 13–19 Nov",
        "W11 20–26 Nov",
        "W12 27 Nov–3 Dec",
        "W13 4–8 Dec",
    ]
    tasks = [
        ("Proposal & planning (both)", "11 Sep – 18 Sep", "Serial start", [1, 1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
        ("Player, camera, shooting (Anthony)", "11 Sep – 25 Sep", "In progress", [1, 1, 1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]),
        ("Wildlife types, rarity, loot (Anthony)", "25 Sep – 22 Oct", "After shooting", [0, 0, 1, 1, 1, 1, 0, 0, 0, 0, 0, 0, 0]),
        ("Health, fire, day/night (Joseph)", "25 Sep – 29 Oct", "Parallel survival", [0, 0, 1, 1, 1, 1, 1, 0, 0, 0, 0, 0, 0]),
        ("Wood, pickaxe, inventory (Joseph)", "2 Oct – 5 Nov", "Needs player", [0, 0, 0, 1, 1, 1, 1, 1, 0, 0, 0, 0, 0]),
        ("Trade post & item prices (both)", "16 Oct – 12 Nov", "Needs inventory", [0, 0, 0, 0, 0, 1, 1, 1, 1, 0, 0, 0, 0]),
        ("Environment, trees, camp (both)", "2 Oct – 12 Nov", "Parallel world", [0, 0, 0, 1, 1, 1, 1, 1, 1, 0, 0, 0, 0]),
        ("Special features: scatter, spoil/cook, legendary, tracks (both)", "16 Oct – 19 Nov", "After loot + fire", [0, 0, 0, 0, 0, 1, 1, 1, 1, 1, 0, 0, 0]),
        ("Playtesting & bug fixes (both)", "30 Oct – 26 Nov", "Ongoing polish", [0, 0, 0, 0, 0, 0, 0, 1, 1, 1, 1, 0, 0]),
        ("Documentation & demo (both)", "13 Nov – 3 Dec", "After stable build", [0, 0, 0, 0, 0, 0, 0, 0, 0, 1, 1, 1, 0]),
        ("Final build & submission (both)", "27 Nov – 8 Dec", "Serial finish", [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 1, 1]),
    ]
    table = doc.add_table(rows=1 + len(tasks), cols=2 + len(weeks))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    write_cell(table.rows[0].cells[0], "Task", bold=True, size=8, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    write_cell(table.rows[0].cells[1], "Dates / type", bold=True, size=8, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    for i, w in enumerate(weeks):
        write_cell(table.rows[0].cells[i + 2], w, bold=True, size=6, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    for r, (name, dates, kind, bars) in enumerate(tasks, start=1):
        row_fill = LIGHT if r % 2 else "FFFFFF"
        write_cell(table.rows[r].cells[0], name, size=8, fill=row_fill)
        write_cell(table.rows[r].cells[1], f"{dates}\n({kind})", size=7, fill=row_fill)
        for i, active in enumerate(bars):
            write_cell(
                table.rows[r].cells[i + 2],
                "██" if active else "",
                size=7,
                center=True,
                color=RGBColor(255, 255, 255),
                fill=GANTT_FILL if active else GANTT_EMPTY,
            )


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    set_page(doc)

    add_paragraph(doc, "UNIVERSITY OF BALAMAND", size=11, bold=True, align="center", space_after=0)
    add_paragraph(doc, "School of Arts and Sciences  •  Computer Science", size=11, italic=True, align="center", space_after=0)
    add_paragraph(doc, "CSIS290 — Project Proposal", size=16, bold=True, align="center", space_before=10, space_after=2)
    add_paragraph(doc, "First-Person Wildlife Hunting Game (Unity)", size=18, bold=True, align="center", space_before=0, space_after=2)
    add_paragraph(
        doc,
        "Prepared by Anthony Gerges and Joseph Chamoun  •  11 September 2026",
        size=11,
        italic=True,
        align="center",
        space_after=14,
    )

    heading(doc, 1, "Project Group Member(s)")
    subheading(doc, "Names & IDs")
    add_paragraph(
        doc,
        "This is a two-member group project.",
        align="justify",
    )
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Full name", "Student ID", "Role"]):
        write_cell(table.rows[0].cells[i], h, bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    write_cell(table.rows[1].cells[0], "Anthony Gerges", size=11, center=True, fill=LIGHT)
    write_cell(table.rows[1].cells[1], "a2212825", size=11, center=True, fill=LIGHT)
    write_cell(table.rows[1].cells[2], "Hunting, wildlife, player combat", size=11, center=True, fill=LIGHT)
    write_cell(table.rows[2].cells[0], "Joseph Chamoun", size=11, center=True, fill="FFFFFF")
    write_cell(table.rows[2].cells[1], "[Insert Joseph’s student ID]", size=11, center=True, fill="FFFFFF")
    write_cell(table.rows[2].cells[2], "Survival, inventory, trading, fire", size=11, center=True, fill="FFFFFF")
    add_paragraph(
        doc,
        "GitHub repository: https://github.com/anthonygerges2004/Hunting-game",
        size=11,
        italic=True,
        space_before=8,
    )

    heading(doc, 2, "Project Abstract")
    subheading(doc, "Brief Description (≤ 50 words)")
    abstract = (
        "This Unity first-person wildlife hunting game lets the player hunt, loot meat, chop wood, and trade. "
        "Birds scatter after a shot, meat spoils unless cooked, and one legendary animal is rare. Fire, night "
        "attacks, and a camp trader complete the wildlife loop."
    )
    words = len(abstract.split())
    add_paragraph(doc, abstract, align="justify", italic=True)
    add_paragraph(doc, f"Word count: {words} words.", size=10, italic=True, space_after=4)

    heading(doc, 3, "Introduction")
    subheading(doc, "Project Overview")
    add_paragraph(
        doc,
        "The project is a first-person hunting game set in a wildlife area, not a score-only shooting gallery. "
        "The player lives off the land: hunt animals, take meat and other materials from kills, gather wood, "
        "keep a fire, stay alive, and trade extra goods for better gear. The world has a morning and a night. "
        "Day is safer for hunting and gathering. Night is dangerous: animals can attack if the player has no fire.",
        align="justify",
        first_line=True,
    )
    add_paragraph(
        doc,
        "Wildlife is mixed. Birds are the most common targets and give a small amount of meat. Medium animals "
        "(for example deer) appear less often and give more loot. Large animals are rare but can be found; they "
        "drop more meat and valuable items, and they are a bigger threat. After a kill, the animal is not only "
        "removed for points: the player loots meat and related materials into an inventory.",
        align="justify",
        first_line=True,
    )
    add_paragraph(
        doc,
        "The player can also harvest wood from trees, but only while carrying a pickaxe. Wood is used to make a "
        "fire at any time and can be sold. A trade post (a camp or trader) is a fixed place on the map where the "
        "player exchanges meat, wood, and weapons. The player has a limited life (health). Reaching zero health "
        "ends the hunt. A campfire nearby at night stops animal attacks; without fire, night wildlife can damage "
        "the player.",
        align="justify",
        first_line=True,
    )
    add_paragraph(
        doc,
        "What makes the project special is that the forest reacts. A gunshot or a kill makes birds fly up and "
        "nearby animals run. Meat is not a permanent score: it spoils unless the player cooks it on the fire or "
        "sells it in time. Cooking restores health, so the fire matters in the day as well as at night. One "
        "legendary animal appears rarely (for example at dawn, after a rumor at the trader). Extra systems include "
        "short-lived tracks, overhunting that empties an area, a trader who is also a rival hunter, and quiet "
        "crouch movement so animals hear running and shooting more easily.",
        align="justify",
        first_line=True,
    )
    add_paragraph(
        doc,
        "A Unity 6 prototype already exists: first-person move and look, raycast shooting, placeholder deer, and "
        "a score HUD. The problem is to grow that prototype into a complete wildlife loop (hunt → loot → cook or "
        "trade → gather → fire → survive day and night) plus the special systems above. The proposed solution keeps "
        "the current player and raycast weapon and adds inventory, rarity, day/night, campfire, trader, trees, health, "
        "and the unique hunting rules (alarm, spoilage, legendary animal).",
        align="justify",
        first_line=True,
    )

    heading(doc, 4, "Motivation and Scope")
    subheading(doc, "Motivation")
    add_paragraph(
        doc,
        "We want a hunting game that also feels like living in the wild. Shooting an animal should matter because "
        "it gives meat, not only a number on the screen. Chopping wood, lighting a fire, and waiting for morning "
        "create tension that a timed score hunt does not. Trading gives a reason to store extra meat and wood "
        "instead of ending the session after a few shots. Special rules make the wildlife feel alive: animals "
        "warn each other, meat spoils, and a legendary animal is the hunt everyone talks about.",
        align="justify",
        first_line=True,
    )
    add_paragraph(
        doc,
        "The idea started from the Unity prototype (player, camera, shoot, score) and from survival games where "
        "night, fire, and gathering change how the player moves. CSIS290 is a good course for this work: two "
        "students can split hunting systems from survival systems, use Unity components, and keep the project on GitHub.",
        align="justify",
        first_line=True,
    )
    subheading(doc, "Scope")
    add_paragraph(doc, "In scope (minimum viable product):", bold=True)
    bullet(doc, "First-person movement, mouse look, and hitscan shooting.")
    bullet(doc, "Wildlife types: common birds, less common medium animals, rare large animals.")
    bullet(doc, "Loot from kills: meat and simple materials go into an inventory.")
    bullet(doc, "Chop wood from trees only if the player has a pickaxe.")
    bullet(doc, "Place or light a campfire at any time (using wood).")
    bullet(doc, "Day/night (morning and night) with different lighting.")
    bullet(doc, "Player health (life limit); game over at zero health.")
    bullet(doc, "At night, animals can attack the player if no fire is active nearby.")
    bullet(doc, "One trade place to buy/sell weapons, meat, and wood.")
    bullet(doc, "A HUD for health, time of day, meat freshness, and inventory counts.")
    bullet(doc, "GitHub repository with a Unity .gitignore.")
    add_paragraph(doc, "Special features (what makes this hunt unique):", bold=True, space_before=6)
    bullet(doc, "Animals warn each other: a shot or kill makes birds scatter and nearby deer flee.")
    bullet(doc, "Meat spoils after a short time unless it is cooked on the fire or sold; cooking restores health.")
    bullet(doc, "One legendary animal is very rare (rumor at the trader; may appear at dawn).")
    bullet(doc, "Large animals leave tracks for a short time; the player can follow them instead of a map marker.")
    bullet(doc, "Overhunting: killing too many animals in one area makes that zone quieter and night more dangerous.")
    bullet(doc, "The trader is a rival hunter who also wants the legendary animal if the player is too slow.")
    bullet(doc, "Crouch / quiet movement: running and shooting alert wildlife more than slow movement.")
    add_paragraph(doc, "Assumptions:", bold=True, space_before=6)
    bullet(doc, "Single-player, keyboard and mouse, Unity editor or Windows build.")
    bullet(doc, "Placeholder models (capsules, simple trees) are allowed until better assets are added.")
    bullet(doc, "One compact map with one trader camp, not an open world.")
    bullet(doc, "“Pickaxe” is the required tool for chopping trees in this game (gameplay rule, not realism).")
    add_paragraph(doc, "Out of scope / boundaries:", bold=True, space_before=6)
    bullet(doc, "Graphic gore, realistic butchering, or blood systems.")
    bullet(doc, "Multiplayer, crafting trees beyond fire, farming, or vehicles.")
    bullet(doc, "Online trading, accounts, mobile ports, or huge streaming maps.")
    bullet(doc, "Full hunting-license simulation or realistic ballistics.")

    heading(doc, 5, "Related Work")
    subheading(doc, "Similar Systems")
    add_paragraph(
        doc,
        "Hunting games such as theHunter: Call of the Wild focus on tracking wildlife and taking a shot. Survival "
        "games such as The Forest, Green Hell, and Valheim add gathering, fire, night danger, and a camp. Our "
        "project sits between those two: hunting is the main activity, but meat, wood, fire, health, and a trader "
        "make the wildlife area a place to live, not only a shooting range.",
        align="justify",
        first_line=True,
    )
    add_paragraph(
        doc,
        "Unity FPS samples show move, look, and shoot. We reuse that idea for the hunter, then add loot, inventory, "
        "and a day/night threat that a normal FPS does not have. Compared with Minecraft-style survival, the focus "
        "stays on hunting animals of different rarity (birds vs. rare large animals) and on a single trade post "
        "rather than a full crafting economy. Special hunting rules (alarming birds, spoiling meat, a legendary "
        "target, tracks, and a rival trader) separate this project from a generic FPS or a simple score hunt.",
        align="justify",
        first_line=True,
    )

    heading(doc, 6, "System Architecture")
    subheading(doc, "Basic Architecture")
    add_paragraph(
        doc,
        "The game is one Unity scene: a wildlife map, trees, a trader camp, and a player. The Player object keeps "
        "the Character Controller, camera, PlayerMove, and PlayerShoot from the prototype, and adds PlayerHealth "
        "and Inventory. Animals are typed (Bird, Medium, RareLarge) with spawn weights, loot tables, and a night "
        "attack behaviour when no campfire is active. DayNightCycle drives lighting. Campfire is a placeable object "
        "that consumes wood and creates a safe radius at night. Trees require a pickaxe in inventory before they "
        "give wood. The Trader is a trigger zone with a simple buy/sell UI and can act as a rival for the legendary "
        "animal. Special scripts handle bird scatter, meat timers, cooking, tracks, overhunting, and crouch detection.",
        align="justify",
        first_line=True,
    )
    add_architecture_table(doc)
    add_flow_table(doc)
    add_paragraph(
        doc,
        "Example data flow: PlayerShoot hits a deer → nearby birds scatter and deer flee → Animal.Die() → "
        "Inventory.Add(meat with a spoil timer). If meat is cooked on the campfire, health is restored; if not, "
        "it becomes worthless. A legendary spawn is rare and may be claimed by the rival trader if the player waits "
        "too long. At night, DayNightCycle sets night. If Campfire.IsActive is false, NightThreat damages "
        "PlayerHealth. Chopping a tree fails unless Inventory.Has(pickaxe).",
        align="justify",
        space_before=8,
        first_line=True,
    )

    heading(doc, 7, "Goals and Objectives")
    subheading(doc, "Final Outcome")
    add_paragraph(
        doc,
        "The expected result is a playable first-person wildlife hunting prototype. A classmate should be able to "
        "hunt, loot meat, cook or sell it before it spoils, follow tracks to a legendary animal, cut wood with a "
        "pickaxe, light a fire, survive a night, and trade at camp. Planned functionalities:",
        align="justify",
    )
    bullet(doc, "Control a first-person hunter (walk, look, shoot, lock/unlock cursor).")
    bullet(doc, "Find birds often; find large animals rarely.")
    bullet(doc, "Collect meat and related items from killed animals.")
    bullet(doc, "Collect wood from trees only with a pickaxe.")
    bullet(doc, "Light a fire at any time using wood.")
    bullet(doc, "See morning and night; night animals attack if there is no fire.")
    bullet(doc, "Have a health/life limit; dying ends the run.")
    bullet(doc, "Trade weapons, meat, and wood at one camp.")
    bullet(doc, "Show health, day/night, meat freshness, and inventory on the HUD.")
    bullet(doc, "Keep the project on GitHub for backup and instructor review.")
    subheading(doc, "Special features")
    add_paragraph(
        doc,
        "These systems are what make the project more than a generic shooter. They are part of the planned outcome:",
        align="justify",
    )
    bullet(doc, "Animals warn each other: birds scatter and nearby animals flee after a shot or a kill.")
    bullet(doc, "Meat spoils unless cooked on the fire or sold in time; cooked meat restores health.")
    bullet(doc, "One legendary animal is very rare and is the main “wow” hunt of the demo.")
    bullet(doc, "Tracks: large/legendary animals leave footprints that fade after a short time.")
    bullet(doc, "Overhunting: too many kills in one area reduces spawns and makes night worse.")
    bullet(doc, "Rival trader: the camp trader also hunts the legendary animal if the player delays.")
    bullet(doc, "Stealth: crouch to stay quieter; running and shooting alert wildlife.")
    add_paragraph(
        doc,
        "Success is a stable Play-mode demo of the full loop (hunt, loot, cook or sell, fire, night, trade) plus "
        "at least the three core special features (scatter, spoiling meat, legendary animal), even if models are placeholders.",
        align="justify",
        first_line=True,
    )

    heading(doc, 8, "Individual Tasks (for Group Projects)")
    subheading(doc, "Task Assignment")
    add_paragraph(
        doc,
        "Work is split so hunting/wildlife and survival/trading can be built in parallel after the shared player exists. "
        "Both members test the merged scene and write the final report.",
        align="justify",
    )
    table = doc.add_table(rows=9, cols=3)
    headers = ["Work package", "Owner", "Justification"]
    rows = [
        (
            "Player, camera, shooting, GitHub setup",
            "Anthony Gerges",
            "Prototype already started by Anthony; combat stays with the hunter systems.",
        ),
        (
            "Animal types, rarity, loot, scatter/flee, legendary spawn, tracks, overhunting",
            "Anthony Gerges",
            "Wildlife reaction and rare hunts depend on shooting and spawn rules.",
        ),
        (
            "Health, day/night, campfire, night attacks, cooking, meat spoil timer",
            "Joseph Chamoun",
            "Survival and food timers sit with fire, health, and inventory.",
        ),
        (
            "Pickaxe, trees, wood, inventory UI, crouch stealth check",
            "Joseph Chamoun",
            "Gathering, inventory, and quiet-move detection are survival-side.",
        ),
        (
            "Trade post and rival-trader legendary race",
            "Joseph Chamoun (Anthony helps legendary spawn)",
            "Trading needs inventory; the rival race needs Anthony’s legendary animal.",
        ),
        (
            "Map, trees, camp visuals, HUD polish",
            "Both",
            "Shared world so hunting and survival meet in one scene.",
        ),
        (
            "Special-feature integration and playtest",
            "Both",
            "Scatter, spoil, and legendary must work together in one demo.",
        ),
        (
            "Testing, proposal, demo, final build",
            "Both",
            "Course delivery is shared; each member tests the other’s systems.",
        ),
    ]
    for i, h in enumerate(headers):
        write_cell(table.rows[0].cells[i], h, bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    for r, row in enumerate(rows, start=1):
        fill = LIGHT if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            write_cell(table.rows[r].cells[c], val, size=10, fill=fill)

    heading(doc, 9, "Gantt Chart")
    subheading(doc, "Project Timeline")
    add_gantt(doc)
    add_paragraph(doc, "Dependencies (serial vs. parallel):", bold=True, space_before=10)
    bullet(doc, "Serial: proposal → player/shooting → loot and survival systems → trade (needs inventory) → final submission.")
    bullet(doc, "Parallel after the player walks and shoots: Anthony’s wildlife/rarity/loot with Joseph’s health, fire, and day/night.")
    bullet(doc, "Parallel: environment (trees, camp) with gathering and animal work.")
    bullet(doc, "Serial: trade UI waits until meat, wood, and inventory exist.")
    bullet(doc, "Serial: special features (scatter, spoil, legendary) start after loot and fire exist.")
    bullet(doc, "Serial finish: documentation and the final build wait until hunt + night + trade work in one scene.")
    add_paragraph(
        doc,
        "Start of project work: 11 September 2026. Target end date: 8 December 2026. "
        "Adjust week labels if the instructor’s official deadline differs.",
        size=10,
        italic=True,
    )

    heading(doc, 10, "Future Work")
    subheading(doc, "Possible Extensions")
    add_paragraph(
        doc,
        "If more time were available after the wildlife loop and special features are stable:",
        align="justify",
    )
    bullet(doc, "Replace placeholders with animated birds, deer, legendary animal, plus audio.")
    bullet(doc, "Wind direction for scent (full hunting simulation).")
    bullet(doc, "Visible rifle, limited ammo, and more weapon tiers at the trader.")
    bullet(doc, "Weather, seasons, a larger map, and more than one trade camp.")
    bullet(doc, "A trophy board at camp for the legendary kill.")
    bullet(doc, "Windows executable and a short GitHub README trailer.")

    heading(doc, 11, "Tools and Technologies")
    subheading(doc, "Tools / Technologies List")
    table = doc.add_table(rows=8, cols=2)
    write_cell(table.rows[0].cells[0], "Tool / technology", bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    write_cell(table.rows[0].cells[1], "Purpose in this project", bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    tools = [
        ("Unity 6 (6000.6) with URP", "3D world, lighting for morning/night, physics, Play-mode testing"),
        ("C# scripts (MonoBehaviour)", "Player, animals, inventory, fire, day/night, trader, health"),
        ("Unity Input System", "WASD, mouse look, shoot, crouch, interact (trees, fire, trade, cook)"),
        ("TextMeshPro / uGUI", "Health, inventory, time of day, trade menu"),
        ("Git + GitHub", "Shared version control (Hunting-game repository)"),
        ("Visual Studio / Cursor (optional)", "Edit C#; Unity remains the main editor"),
        ("Windows 10/11", "Development and playtest platform"),
    ]
    for i, (a, b) in enumerate(tools, start=1):
        fill = LIGHT if i % 2 else "FFFFFF"
        write_cell(table.rows[i].cells[0], a, size=10, fill=fill)
        write_cell(table.rows[i].cells[1], b, size=10, fill=fill)

    heading(doc, 12, "Resources")
    subheading(doc, "Critical Resources")
    add_paragraph(
        doc,
        "The project can be completed with two student PCs, Unity, and free assets. No paid dataset is required. "
        "Instructor contact should be taken from the CSIS290 course page.",
        align="justify",
    )
    table = doc.add_table(rows=8, cols=3)
    for i, h in enumerate(["Resource", "Why it is needed", "How to obtain / contact"]):
        write_cell(table.rows[0].cells[i], h, bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    resources = [
        (
            "Anthony Gerges",
            "Hunting, wildlife, player combat",
            "Student ID a2212825  •  GitHub: anthonygerges2004",
        ),
        (
            "Joseph Chamoun",
            "Survival, inventory, fire, trading",
            "Student ID to be inserted  •  add campus email before submit",
        ),
        (
            "Unity Hub + Unity 6 editor",
            "Build and run the 3D project",
            "https://unity.com/download  (student license if required)",
        ),
        (
            "GitHub repository",
            "Shared backup and history",
            "https://github.com/anthonygerges2004/Hunting-game",
        ),
        (
            "Two Windows PCs with mice",
            "Split development and playtests",
            "Personal machines already in use",
        ),
        (
            "Free Unity assets (optional)",
            "Trees, animals, camp, audio",
            "Unity Asset Store free kits; placeholders allowed",
        ),
        (
            "CSIS290 instructor / lab",
            "Requirements, grading, feedback",
            "Official course email and office hours",
        ),
    ]
    for r, row in enumerate(resources, start=1):
        fill = LIGHT if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            write_cell(table.rows[r].cells[c], val, size=10, fill=fill)

    add_paragraph(
        doc,
        "End of proposal. Insert Joseph Chamoun’s student ID before submission, and confirm the official CSIS290 deadline on the Gantt chart.",
        size=10,
        italic=True,
        space_before=16,
        align="center",
    )

    for path in OUTPUTS:
        try:
            doc.save(path)
            print("saved", path)
        except Exception as exc:
            print("failed", path, exc)


if __name__ == "__main__":
    build()
