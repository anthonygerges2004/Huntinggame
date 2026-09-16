"""Generate CSIS290 Wildlife Hunting Game software-project folder (mirrors Patient Tracker layout)."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT = "Wildlife Hunting Game"
COURSE = "CSIS290"
MEMBERS = "Anthony Gerges (a2212825), Joseph Chamoun"
REPO = "https://github.com/anthonygerges2004/Hunting-game"
SCENE = "GetStarted_Scene.unity"

BASE_DIRS = [
    Path(r"d:\downloadss\Software Project") / PROJECT,
    Path(r"C:\Users\antho\senior\docs\Software Project") / PROJECT,
]

PROPOSAL_SRC = Path(r"C:\Users\antho\senior\CSIS290_Hunting_Game_Project_Proposal.docx")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = "1F3A5F"
LIGHT = "E8EEF6"


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_paragraph(doc, text, *, size=12, bold=False, italic=False, align="left", space_after=8, space_before=0, first_line=False):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    if first_line:
        pf.first_line_indent = Cm(1.0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "9AA4B2")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def write_cell(cell, text, *, bold=False, size=10, center=False, color=None, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_border(cell)
    if fill:
        shade_cell(cell, fill)


def heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{number}. {title}")
    set_run_font(run, size=14, bold=True, color=NAVY)


def subheading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_run_font(run, size=12, bold=True, color=NAVY)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=12)


def title_block(doc, part_title: str, subtitle: str = ""):
    set_page(doc)
    add_paragraph(doc, f"{COURSE} – Group Software Project", size=11, italic=True, align="center")
    add_paragraph(doc, part_title, size=16, bold=True, align="center", space_after=4)
    add_paragraph(doc, PROJECT, size=14, bold=True, align="center", space_after=4)
    if subtitle:
        add_paragraph(doc, subtitle, size=11, italic=True, align="center", space_after=12)
    add_paragraph(doc, MEMBERS, size=11, align="center", space_after=16)


def set_page(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)


def add_table(doc, headers, rows, col_centers=None):
    col_centers = col_centers or []
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        write_cell(table.rows[0].cells[i], h, bold=True, size=11, center=True, color=RGBColor(255, 255, 255), fill=ACCENT)
    for r, row in enumerate(rows, start=1):
        fill = LIGHT if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            write_cell(
                table.rows[r].cells[c],
                val,
                size=10,
                fill=fill,
                center=(c in col_centers),
            )
    doc.add_paragraph()


def build_phase2(out: Path):
    doc = Document()
    title_block(doc, "Project Part #2 – Feasibility Study and Methodology")
    subheading(doc, "Part A – Feasibility Study")
    heading(doc, 1, "Technical Feasibility")
    add_paragraph(
        doc,
        "The Wildlife Hunting Game is a Unity 6 first-person prototype using Universal Render Pipeline (URP), "
        "Character Controller movement, raycast shooting, and C# MonoBehaviour scripts. All technologies are "
        "documented and free for student use.",
        align="justify",
        first_line=True,
    )
    add_table(
        doc,
        ["Aspect", "Assessment"],
        [
            ("Development stack", "Unity 6, URP, Input System, C# scripts (PlayerMove, Animal, Inventory, Campfire, Trader, etc.)"),
            ("Hardware", "Standard Windows PC with GPU; no server required for demo"),
            ("Security", "Offline single-player; GitHub private repo for source control"),
            ("Team capability", "Two CS students: hunting/combat vs survival/inventory/scene art integration"),
        ],
    )
    add_paragraph(doc, "Conclusion: The project is technically feasible with the current prototype and GitHub workflow.", align="justify")
    heading(doc, 2, "Economic Feasibility")
    bullet(doc, "Unity Personal (free), Visual Studio / Rider community, Git and GitHub (free).")
    bullet(doc, "Free Asset Store kits (camping props, grass) and placeholders for wildlife.")
    bullet(doc, "Optional paid ChatGPT Plus for AI art; not required for MVP.")
    add_paragraph(doc, "Conclusion: Zero-budget academic delivery is achievable.", align="justify")
    heading(doc, 3, "Operational Feasibility")
    bullet(doc, "Player uses familiar FPS controls (WASD, mouse, shoot, interact keys documented in README).")
    bullet(doc, "Instructor or classmate can open the scene and play without custom server setup.")
    bullet(doc, "HUD shows health, time of day, and inventory counts.")
    add_paragraph(doc, "Conclusion: Demo and grading workflows are straightforward.", align="justify")
    heading(doc, 4, "Legal Feasibility")
    bullet(doc, "Third-party Unity assets used under Asset Store / free licenses; credited in README.")
    bullet(doc, "No collection of personal user data; single-player local game.")
    add_paragraph(doc, "Conclusion: No blocking legal issues for a course prototype.", align="justify")
    heading(doc, 5, "Schedule Feasibility")
    add_table(
        doc,
        ["Phase", "Deliverable", "Timeline"],
        [
            ("Part 1", "Project proposal", "Weeks 1–2"),
            ("Part 2", "Feasibility and methodology", "Weeks 3–4"),
            ("Part 3", "Requirements analysis", "Weeks 5–6"),
            ("Part 4", "System analysis (DFD, use cases)", "Weeks 7–8"),
            ("Part 5", "Implementation and testing", "Weeks 9–12"),
            ("Part 6", "Design documentation and final demo", "Weeks 13–14"),
        ],
    )
    heading(doc, 6, "Resource Feasibility")
    bullet(doc, f"Human: {MEMBERS}.")
    bullet(doc, f"Repository: {REPO}.")
    bullet(doc, "Knowledge: Unity Learn, course lectures, existing FPS starter assets.")
    add_paragraph(doc, "Conclusion: Resources are available.", align="justify")
    subheading(doc, "Part B – Methodology")
    add_paragraph(
        doc,
        "The team uses an Agile-inspired approach within the semester: a working Unity prototype first, then "
        "iterative additions (inventory, fire, day/night, trader, merge teammate environment work). Documentation "
        "phases run in parallel with implementation. Waterfall-style phase documents (this folder) clarify analysis "
        "for grading; development stays iterative to handle merge and asset import risk.",
        align="justify",
        first_line=True,
    )
    add_table(
        doc,
        ["Criterion", "Waterfall", "Prototyping", "Agile (chosen mix)"],
        [
            ("Approach", "Sequential phases", "Early playable model", "Sprints on gameplay features"),
            ("Requirements", "Fixed in Part 3 doc", "Refined by playtests", "Proposal + playtest feedback"),
            ("Best fit", "Documentation structure", "Unity MVP", "Feature merges on GitHub"),
        ],
    )
    add_paragraph(doc, MEMBERS, align="right", space_before=24)
    doc.save(out / "phase 2 feasibility.docx")


def build_phase3(out: Path):
    doc = Document()
    title_block(doc, "Project Part #3 – Requirements Analysis")
    heading(doc, 1, "Functional Requirements")
    reqs = [
        ("FR-01", "The player shall move in first person (walk, look, optional crouch)."),
        ("FR-02", "The player shall shoot wildlife using raycast hit detection."),
        ("FR-03", "Each animal shall have Species, Size, and Rarity attributes that define loot and spawn behavior."),
        ("FR-04", "Killed animals shall drop meat into the player inventory."),
        ("FR-05", "Meat shall spoil after a time unless cooked on a campfire or sold."),
        ("FR-06", "The player shall chop trees for wood only when a pickaxe is in inventory."),
        ("FR-07", "The player shall light or maintain a campfire using wood."),
        ("FR-08", "The game shall cycle between day and night with visible lighting changes."),
        ("FR-09", "At night, wildlife shall threaten the player if no active campfire is nearby."),
        ("FR-10", "The player shall have health; at zero health the run ends."),
        ("FR-11", "The player shall buy and sell items at a trader camp."),
        ("FR-12", "The HUD shall show health, time of day, and key inventory counts."),
        ("FR-13", "Shots and kills shall alarm nearby wildlife (scatter / flee) where implemented."),
        ("FR-14", "A legendary-rarity hunt may spawn rarely and interact with trader/rival logic where implemented."),
    ]
    add_table(doc, ["ID", "Requirement"], reqs)
    heading(doc, 2, "Non-Functional Requirements")
    add_table(
        doc,
        ["ID", "Requirement"],
        [
            ("NFR-01", "Target playable frame rate on a mid-range student laptop (≥ 30 FPS in main scene)."),
            ("NFR-02", "Single-scene demo load: open Unity project, press Play on main scene."),
            ("NFR-03", "Source stored on GitHub with Unity .gitignore."),
            ("NFR-04", "Controls documented in README for graders."),
            ("NFR-05", "Placeholder or free-store art acceptable for MVP."),
        ],
    )
    heading(doc, 3, "User Classes")
    bullet(doc, "Player (hunter): primary actor for all gameplay.")
    bullet(doc, "Grader / teammate tester: runs build or editor Play mode.")
    bullet(doc, "Trader (in-game NPC): economic interface, not a human login role.")
    add_paragraph(doc, MEMBERS, align="right", space_before=24)
    doc.save(out / "phase 3 requirements.docx")


def build_phase4(out: Path):
    doc = Document()
    title_block(doc, "Project Part #4 – Software Analysis")
    heading(doc, 1, "Context Diagram and Level 0 DFD")
    add_paragraph(
        doc,
        "The Context Diagram (Process 0) and Level 0 DFD are submitted as context diagram.jpeg and level 0.jpeg "
        "in this folder. The following summarizes their structure.",
        align="justify",
        first_line=True,
    )
    subheading(doc, "1.1 Context Diagram – Summary")
    add_table(
        doc,
        ["External entity", "Data flow in", "Data flow out"],
        [
            ("Player", "Movement, aim, shoot, chop, trade, fire commands", "HUD, damage, loot feedback, game over"),
            ("Wildlife", "Spawn and AI behavior", "Meat drops, night attacks, flee/scatter signals"),
            ("Trader camp", "Buy/sell selections", "Items, prices, rumors (legendary)"),
        ],
    )
    subheading(doc, "1.2 Level 0 DFD – Summary")
    add_table(
        doc,
        ["Process", "Name", "Input", "Output"],
        [
            ("1", "Player control & combat", "Input (Player)", "Hit events → P2, damage to Player ← P3"),
            ("2", "Hunt & wildlife", "Hits, spawns", "Loot → D1; alarm → Wildlife"),
            ("3", "Survival & time", "Fire state, time", "Safe night radius; health → D3"),
            ("4", "Gather resources", "Chop action (Player)", "Wood → D1"),
            ("5", "Trade", "Trade keys (Player)", "Inventory updates ↔ D1"),
        ],
        col_centers=[0],
    )
    add_table(
        doc,
        ["Store", "Description"],
        [
            ("D1", "Player inventory – meat (fresh/spoiled/cooked), wood, pickaxe, tools"),
            ("D2", "World session state – time of day, campfire active, overhunt zones"),
            ("D3", "Player vitals – current health and death flag"),
        ],
    )
    heading(doc, 2, "Use Case Scenarios")
    for block in _use_cases():
        subheading(doc, block["title"])
        add_table(doc, ["Field", "Description"], block["rows"])
        add_paragraph(doc, "", space_after=4)
    add_paragraph(doc, MEMBERS, align="right", space_before=24)
    doc.save(out / "phase 4.docx")


def _use_cases():
    def uc(uid, title, desc, actor, trigger, steps, pre, post, fr):
        rows = [
            ("Use Case ID", uid),
            ("Title", title),
            ("Description", desc),
            ("Primary Actor", actor),
            ("Triggering Event", trigger),
            ("Main Success Scenario", steps),
            ("Preconditions", pre),
            ("Postconditions", post),
            ("Requirements Met", fr),
            ("Priority", "High"),
            ("Risk", "Low–Medium"),
        ]
        return {"title": f"{uid}: {title}", "rows": rows}

    return [
        uc(
            "UC-01",
            "Hunt Wildlife",
            "The player aims and shoots an animal. On death, meat is added to inventory and nearby wildlife may flee.",
            "Player",
            "Player fires weapon and raycast hits an Animal.",
            "1. Player aims at wildlife.\n2. Player shoots.\n3. System applies damage.\n4. Animal dies.\n5. Inventory receives meat with spoil timer.\n6. Alarm may scatter birds or flee deer.",
            "Player is in scene with ammo/weapon enabled.",
            "Meat stored in D1; wildlife count reduced.",
            "FR-02, FR-03, FR-04, FR-13",
        ),
        uc(
            "UC-02",
            "Cook or Preserve Meat",
            "The player cooks meat on an active campfire or sells it before spoil time expires.",
            "Player",
            "Player presses cook near fire or sells at trader.",
            "1. Player has raw meat in D1.\n2. Player activates campfire or trader sell.\n3. System validates proximity/stock.\n4. Meat becomes cooked or currency/items update.\n5. HUD refreshes counts.",
            "Raw meat in inventory; campfire lit or trader in range.",
            "Cooked meat restores health when consumed (if implemented) or spoil timer cleared via sale.",
            "FR-05, FR-07, FR-11",
        ),
        uc(
            "UC-03",
            "Survive the Night",
            "During night phase, the player avoids damage by staying near an active campfire.",
            "Player",
            "DayNightCycle enters night.",
            "1. Lighting switches to night.\n2. If campfire inactive, night threat damages PlayerHealth.\n3. If campfire active in range, player remains safe.\n4. HUD shows time of day.",
            "Player alive in scene.",
            "Player reaches morning or dies (health in D3).",
            "FR-08, FR-09, FR-10",
        ),
        uc(
            "UC-04",
            "Gather Wood",
            "The player chops a tree when carrying a pickaxe and receives wood in inventory.",
            "Player",
            "Player presses chop near HarvestableTree.",
            "1. System checks pickaxe in D1.\n2. If present, wood added to D1.\n3. If absent, feedback shown.",
            "Player near tree.",
            "Wood count in D1 increased.",
            "FR-06",
        ),
        uc(
            "UC-05",
            "Trade at Camp",
            "The player buys a pickaxe or sells meat/wood at the trader trigger.",
            "Player",
            "Player enters trader zone and presses trade hotkeys.",
            "1. Trader UI or key handler activates.\n2. System validates inventory.\n3. Items and counts update in D1.\n4. HUD updates.",
            "Player in trader trigger.",
            "Inventory reflects sale/purchase.",
            "FR-11",
        ),
    ]


def build_phase5(out: Path):
    doc = Document()
    title_block(doc, "Project Part #5 – Implementation and Testing")
    heading(doc, 1, "Implementation Summary")
    add_paragraph(
        doc,
        f"Main scene: {SCENE}. Core scripts under Assets/: PlayerMove, PlayerShoot, PlayerHealth, Animal, "
        "Inventory, GameHUD, DayNightCycle, Campfire, HarvestableTree, Trader, WildlifeAlarm, Overhunt. "
        "Environment and props integrated via GitHub (grass, URP camping props, campfire/tree fixes).",
        align="justify",
        first_line=True,
    )
    add_table(
        doc,
        ["Component", "Script(s)", "Status"],
        [
            ("First-person control", "PlayerMove, Character Controller", "Implemented"),
            ("Combat", "PlayerShoot", "Implemented"),
            ("Wildlife", "Animal (+ Species/Size/Rarity fields)", "Implemented / data-driven"),
            ("Inventory & HUD", "Inventory, GameHUD", "Implemented"),
            ("Survival", "Campfire, DayNightCycle, PlayerHealth", "Implemented"),
            ("Gathering", "HarvestableTree", "Implemented"),
            ("Trading", "Trader", "Implemented"),
            ("Special rules", "WildlifeAlarm, Overhunt", "Implemented / partial polish"),
        ],
    )
    heading(doc, 2, "Controls (demo)")
    add_table(
        doc,
        ["Input", "Action"],
        [
            ("WASD + mouse", "Move and look"),
            ("Left click", "Shoot"),
            ("Ctrl", "Crouch / stealth"),
            ("F", "Campfire"),
            ("C", "Cook"),
            ("E", "Chop tree"),
            ("1 / 3", "Sell meat / buy pickaxe at trader"),
        ],
    )
    heading(doc, 3, "System Testing Checklist")
    tests = [
        ("ST-01", "Play scene 5 minutes without console errors", "Pass / Fail"),
        ("ST-02", "Kill deer/bird → meat in HUD", "Pass / Fail"),
        ("ST-03", "Night without fire → health drops", "Pass / Fail"),
        ("ST-04", "Fire lit → night safe in radius", "Pass / Fail"),
        ("ST-05", "Chop without pickaxe blocked", "Pass / Fail"),
        ("ST-06", "Trader buy pickaxe then chop works", "Pass / Fail"),
    ]
    add_table(doc, ["ID", "Test", "Result (fill on demo day)"], tests)
    add_paragraph(
        doc,
        "Record pass/fail and screenshot Play mode for submission. Export Windows build if required by syllabus.",
        italic=True,
    )
    add_paragraph(doc, MEMBERS, align="right", space_before=24)
    doc.save(out / "phase 5 implementation.docx")


def build_phase6(out: Path):
    doc = Document()
    title_block(doc, "Phase 6 – Design Documentation (Part B)", PROJECT)
    heading(doc, 3, "Data Design")
    add_paragraph(
        doc,
        "Runtime data is held in Unity components and serializable fields rather than a SQL database. Logical data "
        "stores match DFD D1–D3.",
        align="justify",
        first_line=True,
    )
    bullet(doc, "D1 Inventory: meat stacks with spoil timestamps, wood count, pickaxe flag, cooked meat.")
    bullet(doc, "D2 World state: DayNightCycle phase, Campfire.IsActive, Overhunt region counters.")
    bullet(doc, "D3 Player vitals: PlayerHealth current/max, death event.")
    bullet(doc, "Wildlife template: Species (Bird, Deer, Bear…), Size (Small/Medium/Large), Rarity (Common→Legendary).")
    subheading(doc, "Relationships")
    bullet(doc, "One Player → one Inventory and one PlayerHealth.")
    bullet(doc, "Many Animal instances; each references wildlife attribute triple.")
    bullet(doc, "Trader reads/writes Inventory only when player in trigger.")
    heading(doc, 4, "Reusability")
    bullet(doc, "PlayerMove / PlayerShoot reused from FPS starter pattern.")
    bullet(doc, "Inventory API (Add, Has, counts) reused by Campfire, Trader, HarvestableTree, Animal.")
    bullet(doc, "GameHUD listens to Inventory and DayNightCycle for display.")
    bullet(doc, "URP assets and camping prop prefabs reused across scene.")
    heading(doc, 5, "Software Risks & Mitigation")
    add_table(
        doc,
        ["Risk", "Probability", "Impact", "Mitigation"],
        [
            ("Git merge conflicts on scene", "Medium", "High", "Pull before opening Unity; communicate; scene owner"),
            ("Large package import time", "Medium", "Medium", "Document packages; optional trim AI packages"),
            ("Asset inconsistency", "Medium", "Low", "Free kits + placeholders; AI icons optional"),
            ("Scope creep on special features", "Medium", "Medium", "MVP list in Part 3; polish if time"),
            ("Performance on low GPU", "Low", "Medium", "URP quality settings; reduce grass density"),
        ],
    )
    heading(doc, 6, "Software Testing Plan")
    add_table(
        doc,
        ["Level", "Purpose", "Example", "Responsible"],
        [
            ("Unit", "Script logic", "Inventory.Add, spoil timer", "Developer"),
            ("Integration", "Systems together", "Shoot → loot → HUD", "Both"),
            ("System", "Full loop", "Hunt → cook → night → trade", "Both"),
            ("Acceptance", "Course requirements", "Grader checklist Part 5", "Instructor"),
            ("Regression", "After merge", "Re-run ST-01–ST-06", "Both"),
        ],
    )
    add_paragraph(
        doc,
        "Conclusion: Composable wildlife attributes, layered Unity architecture, and the Part 5 checklist support "
        "a maintainable hunting-survival demo aligned with Parts 3–5.",
        align="justify",
        first_line=True,
    )
    add_paragraph(doc, MEMBERS, align="right", space_before=24)
    doc.save(out / "phase 6.docx")
    shutil.copy(out / "phase 6.docx", out / "Wildlife_Hunting_Game_Phase_6_Design_Documentation.docx")


def write_readme(out: Path):
    text = f"""Wildlife Hunting Game – Software Project Package
{COURSE} | {MEMBERS}
Repo: {REPO}

Contents (mirror Patient Tracker Software Project folder):
  - CSIS290_Hunting_Game_Project_Proposal.docx  (Part 1 – copy from project root)
  - phase 2 feasibility.docx                      Part 2
  - phase 3 requirements.docx                     Part 3
  - phase 4.docx                                  Part 4 (see also diagrams)
  - phase 5 implementation.docx                   Part 5
  - phase 6.docx                                  Part 6
  - Wildlife_Hunting_Game_Phase_6_Design_Documentation.docx  (same as phase 6)
  - context diagram.jpeg
  - level 0.jpeg

Export Part 1 and phases to PDF for submission if required.
Fill Pass/Fail in phase 5 after final playtest.
Insert Joseph's student ID on proposal if still placeholder.
"""
    (out / "README.txt").write_text(text, encoding="utf-8")


def copy_proposal(out: Path):
    if PROPOSAL_SRC.is_file():
        shutil.copy(PROPOSAL_SRC, out / "CSIS290_Hunting_Game_Project_Proposal.docx")
    else:
        (out / "MISSING_PROPOSAL.txt").write_text(
            f"Copy proposal manually from {PROPOSAL_SRC}", encoding="utf-8"
        )


def build_all():
    for base in BASE_DIRS:
        base.mkdir(parents=True, exist_ok=True)
        build_phase2(base)
        build_phase3(base)
        build_phase4(base)
        build_phase5(base)
        build_phase6(base)
        copy_proposal(base)
        write_readme(base)
        print("built", base)


if __name__ == "__main__":
    build_all()
